import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_docx_report(output_docx_path: str):
    doc = Document()

    # Configurar margenes de pagina (1 pulgada = 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Colores
    c_primary = RGBColor(13, 71, 161)     # #0d47a1
    c_secondary = RGBColor(21, 101, 192)  # #1565c0
    c_dark = RGBColor(33, 33, 33)         # #212121
    c_gray = RGBColor(100, 116, 139)

    # Estilos de encabezados
    # Title
    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_u1 = p_univ.add_run("INSTITUCIÓN UNIVERSITARIA DIGITAL DE ANTIOQUIA\n")
    r_u1.bold = True
    r_u1.font.size = Pt(13)
    r_u1.font.color.rgb = RGBColor(71, 85, 105)
    r_u2 = p_univ.add_run("FACULTAD DE INGENIERÍA | INGENIERÍA DE SOFTWARE")
    r_u2.font.size = Pt(10)
    r_u2.font.color.rgb = c_gray

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(20)

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("DOCUMENTO DE ARQUITECTURA DE BIG DATA Y MODELO DE DATOS\n")
    r_t.bold = True
    r_t.font.size = Pt(22)
    r_t.font.color.rgb = c_primary

    r_st = p_title.add_run("Actividad Evaluativa 4 (EA4) - Proyecto Integrador")
    r_st.font.size = Pt(13)
    r_st.font.color.rgb = c_secondary

    doc.add_paragraph().paragraph_format.space_before = Pt(25)

    # Tabla de metadatos de portada
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_meta.autofit = False

    meta_rows = [
        ("Asignatura:", "Arquitectura Big Data (7° Semestre)"),
        ("Docente Asesor:", "Facultad de Ingeniería - IU Digital de Antioquia"),
        ("Pipeline Implementado:", "Ingestión API (EA1) ➔ Limpieza (EA2) ➔ Enriquecimiento Multiformato (EA3) ➔ Arquitectura (EA4)"),
        ("Dataset Base:", "FIFA 20 Complete Player Dataset (18,278 registros) + 6 Fuentes Multiformato"),
        ("Entorno de Ejecución:", "Simulación Cloud (SQLite Data Lakehouse) + GitHub Actions CI/CD Automatizado"),
        ("Fecha:", "Septiembre 2026")
    ]

    for i, (k, v) in enumerate(meta_rows):
        row = table_meta.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.width = Inches(2.0)
        c1.width = Inches(4.5)
        
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(k)
        r0.bold = True
        r0.font.size = Pt(9.5)
        r0.font.color.rgb = c_primary

        p1 = c1.paragraphs[0]
        r1 = p1.add_run(v)
        r1.font.size = Pt(9.5)
        r1.font.color.rgb = c_dark

        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F8FAFC")
        set_cell_margins(c0, top=120, bottom=120, left=150, right=150)
        set_cell_margins(c1, top=120, bottom=120, left=150, right=150)

    doc.add_page_break()

    # =========================================================================
    # SECCIÓN 1
    # =========================================================================
    h1 = doc.add_heading(level=1)
    r = h1.add_run("1. Introducción y Descripción Global de la Arquitectura")
    r.font.color.rgb = c_primary
    r.bold = True

    p = doc.add_paragraph(
        "El presente documento técnico detalla la arquitectura de software y datos diseñada e implementada para el "
        "Proyecto Integrador de Big Data en la Institución Universitaria Digital de Antioquia. El sistema cubre de forma "
        "exhaustiva el ciclo de vida integral del dato, abarcando las cuatro etapas metodológicas establecidas:"
    )
    p.runs[0].font.color.rgb = c_dark
    p.runs[0].font.size = Pt(10.5)

    phases = [
        ("Fase 1 (EA1 - Ingestión):", " Conexión automatizada vía HTTP a servicios web y APIs REST públicas, extrayendo datos semiestructurados (JSON) y persistiendo su estado crudo en una base de datos relacional analítica."),
        ("Fase 2 (EA2 - Preprocesamiento y Limpieza):", " Depuración masiva de inconsistencias, eliminación de duplicados, imputación estadística de valores nulos (mediana e imputaciones por dominio), corrección de tipos e ingeniería de variables biométricas."),
        ("Fase 3 (EA3 - Enriquecimiento Multiformato):", " Integración sinérgica de datos complementarios provenientes de 6 formatos heterogéneos (JSON, CSV, XML, HTML, TXT y XLSX) mediante operaciones relacionales de cruce (LEFT JOIN) preservando la integridad referencial al 100%."),
        ("Fase 4 (EA4 - Modelo de Datos y Automatización CI/CD):", " Consolidación del esquema dimensional analítico en SQLite, generación de evidencias ejecutivas y orquestación desatendida mediante GitHub Actions.")
    ]
    for bold_prefix, text in phases:
        bp = doc.add_paragraph(style='List Bullet')
        r_b = bp.add_run(bold_prefix)
        r_b.bold = True
        r_b.font.color.rgb = c_primary
        r_t = bp.add_run(text)
        r_t.font.color.rgb = c_dark
        r_b.font.size = Pt(10)
        r_t.font.size = Pt(10)

    p_med = doc.add_paragraph(
        "La solución adopta el patrón de Arquitectura Medallion (Data Lakehouse), dividiendo el flujo en capas de madurez: "
        "Bronze (Datos Crudos / Ingestión), Silver (Datos Depurados / Limpieza) y Gold (Datos Enriquecidos / Analítica y Modelado), "
        "garantizando trazabilidad, reproducibilidad y gobernanza en todo momento."
    )
    p_med.runs[0].font.color.rgb = c_dark
    p_med.runs[0].font.size = Pt(10.5)

    # =========================================================================
    # SECCIÓN 2
    # =========================================================================
    h2 = doc.add_heading(level=1)
    r = h2.add_run("2. Diagrama de Arquitectura y Flujo de Procesamiento")
    r.font.color.rgb = c_primary
    r.bold = True

    p = doc.add_paragraph(
        "El siguiente diagrama ilustra la interacción entre las fuentes heterogéneas, los scripts de procesamiento en Python, "
        "la persistencia escalonada en SQLite y el pipeline de automatización continua CI/CD en GitHub Actions:"
    )
    p.runs[0].font.color.rgb = c_dark

    diag_arq_path = os.path.join(os.path.dirname(__file__), "diagrama_arquitectura.png")
    if os.path.exists(diag_arq_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(diag_arq_path, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figura 1: Arquitectura Medallion End-to-End y Pipeline CI/CD en GitHub Actions.")
        r_cap.italic = True
        r_cap.font.size = Pt(9)
        r_cap.font.color.rgb = c_gray

    # =========================================================================
    # SECCIÓN 3
    # =========================================================================
    doc.add_page_break()
    h3 = doc.add_heading(level=1)
    r = h3.add_run("3. Modelo de Datos y Diagrama Entidad - Relación (ER)")
    r.font.color.rgb = c_primary
    r.bold = True

    p = doc.add_paragraph(
        "El modelo de datos se estructura bajo un enfoque analítico optimizado para consultas de alto rendimiento y modelos de "
        "Machine Learning. La entidad central enriched_players combina los atributos intrínsecos de los jugadores con las dimensiones externas integradas:"
    )
    p.runs[0].font.color.rgb = c_dark

    diag_er_path = os.path.join(os.path.dirname(__file__), "diagrama_modelo_er.png")
    if os.path.exists(diag_er_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.add_run().add_picture(diag_er_path, width=Inches(6.2))
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_cap = p_cap.add_run("Figura 2: Diagrama Entidad-Relación (ER) del Modelo de Datos Enriquecido.")
        r_cap.italic = True
        r_cap.font.size = Pt(9)
        r_cap.font.color.rgb = c_gray

    h3_sub = doc.add_heading(level=2)
    r_sub = h3_sub.add_run("Diccionario de Datos y Estructura de Tablas:")
    r_sub.font.color.rgb = c_secondary

    dict_rows = [
        ("Tabla / Entidad", "Formato", "Clave (PK / FK)", "Atributos Relevantes"),
        ("enriched_players (Tabla Maestra Gold)", "SQLite DB", "PK: sofifa_id\nFK: nationality, club", "sofifa_id, short_name, overall, potential, value_eur, wage_eur, bmi, potential_growth, star_rating_index (44 cols)"),
        ("countries_info", "JSON (.json)", "PK: nationality", "continent, fifa_confederation, country_iso3, fifa_ranking_tier"),
        ("club_leagues", "CSV (.csv)", "PK: club", "league_name, league_country, league_tier, is_top_5_european_league"),
        ("stadiums_info", "XML (.xml)", "PK: club", "stadium_name, stadium_capacity"),
        ("national_trophies", "HTML (.html)", "PK: nationality", "world_cup_titles, continental_trophies"),
        ("player_contracts", "TXT (.txt)", "PK: sofifa_id", "contract_tier, rep_stars"),
        ("sponsorship_tiers", "XLSX (.xlsx)", "PK: club", "main_sponsor, sponsor_tier")
    ]

    t_dict = doc.add_table(rows=len(dict_rows), cols=4)
    t_dict.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_dict.autofit = False

    col_w = [Inches(1.8), Inches(1.0), Inches(1.4), Inches(2.3)]
    for r_idx, r_data in enumerate(dict_rows):
        row = t_dict.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = col_w[c_idx]
            p_c = cell.paragraphs[0]
            r_c = p_c.add_run(val)
            if r_idx == 0:
                set_cell_background(cell, "0D47A1")
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(255, 255, 255)
                r_c.font.size = Pt(9)
            else:
                bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                r_c.font.color.rgb = c_dark
                r_c.font.size = Pt(8.5)
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    p_just = doc.add_paragraph()
    p_just.paragraph_format.space_before = Pt(10)
    r_jb = p_just.add_run("Justificación del Modelo: ")
    r_jb.bold = True
    r_jb.font.color.rgb = c_secondary
    r_jt = p_just.add_run(
        "Se adoptó un modelo híbrido relacional/desnormalizado tipo Estrella (Star Schema). Esta estructura permite entrenar "
        "algoritmos de agrupamiento (Clustering), clasificación y regresión sin incurrir en penalizaciones computacionales "
        "por múltiples uniones en tiempo de ejecución, optimizando los tiempos de procesamiento en Big Data."
    )
    r_jt.font.color.rgb = c_dark

    # =========================================================================
    # SECCIÓN 4
    # =========================================================================
    doc.add_page_break()
    h4 = doc.add_heading(level=1)
    r = h4.add_run("4. Justificación de Herramientas y Tecnologías")
    r.font.color.rgb = c_primary
    r.bold = True

    tools = [
        ("SQLite3 (Motor de Persistencia Analítica):", " Actúa como emulador local de alta eficiencia de un Data Lakehouse / Data Warehouse en la nube. Proporciona almacenamiento ligero, embebido, sin requerir despliegue de infraestructura pesada en etapas de desarrollo, ofreciendo compatibilidad SQL estándar y transacciones ACID."),
        ("Pandas y NumPy (Capa de Preprocesamiento y ETL):", " Librerías líderes en el ecosistema de Python para manipulación vectorial en memoria. Permiten transformaciones matriciales, tratamiento masivo de valores nulos, normalizaciones Min-Max y cálculos vectorizados (como el BMI e índices compuestos) con alta velocidad."),
        ("PySpark / Ecosistema Distribuido (Simulación de Escalabilidad):", " Para cargas de trabajo que superan la memoria RAM de un nodo (escala de Terabytes), el pipeline está concebido para interoperar con PySpark mediante DataFrames distribuidos (Resilient Distributed Datasets / Catalyst Optimizer) en clusters de procesamiento."),
        ("GitHub Actions (Orquestación CI/CD y Calidad de Datos):", " Plataforma de automatización de flujos de trabajo en la nube. Permite ejecutar de forma desatendida el pipeline ante cada cambio en el código, validar aserciones de integridad de datos mediante pruebas automáticas y publicar artefactos de evidencia (bases de datos, hojas Excel y reportes TXT).")
    ]
    for t_name, t_desc in tools:
        bp = doc.add_paragraph(style='List Bullet')
        r_b = bp.add_run(t_name)
        r_b.bold = True
        r_b.font.color.rgb = c_primary
        r_t = bp.add_run(t_desc)
        r_t.font.color.rgb = c_dark
        r_b.font.size = Pt(10)
        r_t.font.size = Pt(10)

    # =========================================================================
    # SECCIÓN 5
    # =========================================================================
    h5 = doc.add_heading(level=1)
    r = h5.add_run("5. Simulación del Entorno Cloud vs Nube Real")
    r.font.color.rgb = c_primary
    r.bold = True

    p = doc.add_paragraph(
        "El proyecto fue diseñado bajo el principio de Portabilidad Cloud-Native. La siguiente tabla presenta el mapeo directo "
        "de cada componente local simulado hacia los servicios administrados de los tres proveedores líderes de nube (AWS, GCP y Azure):"
    )
    p.runs[0].font.color.rgb = c_dark

    cloud_rows = [
        ("Fase / Componente", "Simulación Local", "AWS", "GCP", "Microsoft Azure"),
        ("Almacenamiento Crudo (Bronze)", "players_20.csv / SQLite", "Amazon S3 (Raw Bucket)", "Google Cloud Storage (GCS)", "Azure Blob / ADLS Gen2"),
        ("Motor de Procesamiento (Silver)", "Python / Pandas", "AWS Glue / EMR Spark", "Dataproc / Cloud Dataflow", "Azure Synapse Spark Pool"),
        ("Capa Analítica (Gold)", "SQLite (ingestion.db)", "Amazon Athena / Redshift", "Google BigQuery", "Azure Synapse Analytics"),
        ("Orquestación y CI/CD", "GitHub Actions", "AWS Step Functions / MWAA", "Cloud Composer (Airflow)", "Azure Data Factory (ADF)")
    ]

    t_cloud = doc.add_table(rows=len(cloud_rows), cols=5)
    t_cloud.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_cloud.autofit = False

    c_w = [Inches(1.4), Inches(1.2), Inches(1.3), Inches(1.3), Inches(1.3)]
    for r_idx, r_data in enumerate(cloud_rows):
        row = t_cloud.rows[r_idx]
        for c_idx, val in enumerate(r_data):
            cell = row.cells[c_idx]
            cell.width = c_w[c_idx]
            p_c = cell.paragraphs[0]
            r_c = p_c.add_run(val)
            if r_idx == 0:
                set_cell_background(cell, "1565C0")
                r_c.bold = True
                r_c.font.color.rgb = RGBColor(255, 255, 255)
                r_c.font.size = Pt(8.5)
            else:
                bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                set_cell_background(cell, bg)
                r_c.font.color.rgb = c_dark
                r_c.font.size = Pt(8)
            set_cell_margins(cell, top=80, bottom=80, left=90, right=90)

    # =========================================================================
    # SECCIÓN 6
    # =========================================================================
    doc.add_page_break()
    h6 = doc.add_heading(level=1)
    r = h6.add_run("6. Flujo de Datos y Automatización en GitHub Actions")
    r.font.color.rgb = c_primary
    r.bold = True

    p = doc.add_paragraph(
        "El flujo de trabajo automatizado (configurado en .github/workflows/bigdata.yml) asegura la reproducibilidad técnica completa. "
        "En cada ejecución se siguen las siguientes fases:"
    )
    p.runs[0].font.color.rgb = c_dark

    ci_steps = [
        ("1. Aprovisionamiento:", " Instancia virtual Ubuntu Latest con Python 3.11 y caché de dependencias."),
        ("2. Instalación de Dependencias:", " Instalación determinística desde requirements.txt."),
        ("3. Ejecución del Pipeline:", " Invocación secuencial de la ingesta, preprocesamiento y enriquecimiento multiformato."),
        ("4. Pruebas de Aserción e Integridad:", " Validación formal de la existencia no vacía de ingestion.db, enriched_data.xlsx y enriched_report.txt."),
        ("5. Auditoría en Logs:", " Despliegue en tiempo real del informe de auditoría con las tasas de coincidencia del 100%."),
        ("6. Publicación de Artefactos:", " Empaquetado y almacenamiento de las evidencias (retención de 30 días) accesibles para revisión docente.")
    ]
    for s_title, s_desc in ci_steps:
        bp = doc.add_paragraph(style='List Bullet')
        r_b = bp.add_run(s_title)
        r_b.bold = True
        r_b.font.color.rgb = c_primary
        r_t = bp.add_run(s_desc)
        r_t.font.color.rgb = c_dark
        r_b.font.size = Pt(9.5)
        r_t.font.size = Pt(9.5)

    # =========================================================================
    # SECCIÓN 7
    # =========================================================================
    h7 = doc.add_heading(level=1)
    r = h7.add_run("7. Conclusiones y Recomendaciones")
    r.font.color.rgb = c_primary
    r.bold = True

    h7_b = doc.add_heading(level=2)
    r = h7_b.add_run("Principales Beneficios de la Arquitectura:")
    r.font.color.rgb = c_secondary

    benefits = [
        ("Integridad y Calidad Garantizada:", " Reducción del 100% de valores nulos en variables clave y cero pérdida de registros en las operaciones de cruce multiformato."),
        ("Modularidad y Mantenibilidad:", " Separación clara de responsabilidades en scripts independientes para ingesta, limpieza y enriquecimiento."),
        ("Automatización Continua:", " Eliminación de tareas operativas manuales gracias a GitHub Actions.")
    ]
    for b_title, b_desc in benefits:
        bp = doc.add_paragraph(style='List Bullet')
        r_b = bp.add_run(b_title)
        r_b.bold = True
        r_b.font.color.rgb = c_secondary
        r_t = bp.add_run(b_desc)
        r_t.font.color.rgb = c_dark

    h7_r = doc.add_heading(level=2)
    r = h7_r.add_run("Recomendaciones para Implementación en Producción:")
    r.font.color.rgb = c_secondary

    recs = [
        ("Adopción de Almacenamiento Formato Parquet / Delta Lake:", " Migrar el formato de almacenamiento de SQLite a Apache Parquet con compresión Snappy para optimizar lecturas columnares en Big Data."),
        ("Implementación de Data Governance:", " Incorporar herramientas de catálogo de datos como Apache Atlas o Google Dataplex para trazabilidad de linaje."),
        ("Monitoreo y Alertas en Tiempo Real:", " Integrar alertas automáticas vía Slack o correo ante caídas en las tasas de match o fallos en las fuentes externas.")
    ]
    for r_title, r_desc in recs:
        bp = doc.add_paragraph(style='List Bullet')
        r_b = bp.add_run(r_title)
        r_b.bold = True
        r_b.font.color.rgb = c_secondary
        r_t = bp.add_run(r_desc)
        r_t.font.color.rgb = c_dark

    p_end = doc.add_paragraph()
    p_end.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_end.paragraph_format.space_before = Pt(20)
    r_end = p_end.add_run("Fin del Documento Técnico de Arquitectura | Institución Universitaria Digital de Antioquia")
    r_end.italic = True
    r_end.font.size = Pt(8.5)
    r_end.font.color.rgb = c_gray

    doc.save(output_docx_path)
    print(f"[OK] Documento Word (.docx) generado exitosamente en: {output_docx_path}")


if __name__ == "__main__":
    docs_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs"))
    os.makedirs(docs_dir, exist_ok=True)
    docx_target = os.path.join(docs_dir, "arquitectura_modelo.docx")
    create_docx_report(docx_target)
